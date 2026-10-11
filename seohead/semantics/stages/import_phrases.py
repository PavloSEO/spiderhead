"""Preview and import CSV/XLSX phrase data with explicit field mappings."""

from __future__ import annotations

import csv
import json
from datetime import datetime, timezone
from pathlib import Path

from seohead.semantics.demand import has_observed_demand
from seohead.semantics.norm import normalize

_FIELDS = ("query", "frequency", "group", "subgroup", "url", "comment")
MAX_IMPORT_BYTES = 32 * 1024 * 1024
MAX_IMPORT_ROWS = 100_000
MAX_IMPORT_COLUMNS = 128
MAX_IMPORT_CELL_BYTES = 16_384
MAX_IMPORT_TEXT_BYTES = 16 * 1024 * 1024
_DEFAULTS = {
    "query": ("norm", "phrase", "query", "фраза", "запрос"),
    "frequency": ("frequency", "base", "частотность", "частота"),
    "group": ("group", "category", "категория", "группа", "cluster"),
    "subgroup": ("subgroup", "подгруппа"),
    "url": ("url", "landing", "landing_url", "страница"),
    "comment": ("comment", "комментарий", "примечание"),
    "impr": ("impr", "impressions", "показы"),
    "pos_y": ("pos_y", "yandex_pos", "позиция_яндекс"),
    "pos_g": ("pos_g", "google_pos", "позиция_google"),
}


def _collect_rows(source_rows, headers, *, source, strict_width):
    if not 1 <= len(headers) <= MAX_IMPORT_COLUMNS:
        raise ValueError(f"import must have between 1 and {MAX_IMPORT_COLUMNS} columns")
    rows, text_bytes = [], 0
    for header in headers:
        size = len(header.encode("utf-8"))
        if size > MAX_IMPORT_CELL_BYTES:
            raise ValueError(f"column header exceeds the {MAX_IMPORT_CELL_BYTES}-byte cell limit")
        text_bytes += size
    if text_bytes > MAX_IMPORT_TEXT_BYTES:
        raise ValueError(f"import exceeds the {MAX_IMPORT_TEXT_BYTES}-byte text limit")
    for line, raw in enumerate(source_rows, 2):
        if line > MAX_IMPORT_ROWS + 1:
            raise ValueError(f"import exceeds the {MAX_IMPORT_ROWS}-row limit")
        if isinstance(raw, dict):
            if None in raw or (strict_width and any(raw.get(header) is None for header in headers)):
                raise ValueError(f"wrong column count on {source} row {line}")
            values = [raw.get(header) for header in headers]
        else:
            if len(raw) > len(headers):
                raise ValueError(f"wrong column count on {source} row {line}")
            values = list(raw) + [None] * (len(headers) - len(raw))
        if any(
            not header and value not in (None, "")
            for header, value in zip(headers, values, strict=True)
        ):
            raise ValueError(f"non-empty cell has no column header on {source} row {line}")
        for value in values:
            if value is None:
                continue
            size = len(str(value).encode("utf-8"))
            if size > MAX_IMPORT_CELL_BYTES:
                raise ValueError(
                    f"cell exceeds the {MAX_IMPORT_CELL_BYTES}-byte limit on {source} row {line}"
                )
            text_bytes += size
            if text_bytes > MAX_IMPORT_TEXT_BYTES:
                raise ValueError(f"import exceeds the {MAX_IMPORT_TEXT_BYTES}-byte text limit")
        rows.append(dict(zip(headers, values, strict=True)))
    return rows


def _read_table(path: Path, sheet: str | None):
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet:
                if sheet not in workbook.sheetnames:
                    raise ValueError(f"worksheet not found: {sheet}")
                ws = workbook[sheet]
            else:
                ws = workbook.active
            values = ws.iter_rows(values_only=True)
            headers = next(values, None)
            if headers is None:
                raise ValueError("worksheet needs a header row")
            headers = [str(h).strip() if h is not None else "" for h in headers]
            if len({h.casefold() for h in headers if h}) != len([h for h in headers if h]):
                raise ValueError("worksheet headers must be unique")
            rows = _collect_rows(
                values, headers, source=f"worksheet {ws.title}", strict_width=False
            )
            return headers, rows
        finally:
            workbook.close()
    if sheet:
        raise ValueError("--sheet is available only for .xlsx and .xlsm imports")
    with path.open(encoding="utf-8-sig", newline="") as stream:
        sample = stream.read(4096)
        stream.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(stream, dialect=dialect, strict=True)
        if not reader.fieldnames:
            raise ValueError("import file needs a header row")
        headers = [header.strip() if header else "" for header in reader.fieldnames]
        if len({h.casefold() for h in headers if h}) != len([h for h in headers if h]):
            raise ValueError("CSV headers must be unique")
        reader.fieldnames = headers
        try:
            rows = _collect_rows(reader, headers, source="CSV", strict_width=True)
        except csv.Error as error:
            raise ValueError(f"malformed CSV: {error}") from error
        return headers, rows


def _mapping(headers, mapping=None):
    mapping = dict(mapping or {})
    unknown = set(mapping) - set(_FIELDS) - {"base", "quoted", "exact", "impr", "pos_y", "pos_g"}
    if unknown:
        raise ValueError(f"unsupported import mapping fields: {', '.join(sorted(unknown))}")
    result = {}
    lower = {header.casefold(): header for header in headers if header}
    for field in (*_FIELDS, "quoted", "exact", "impr", "pos_y", "pos_g"):
        requested = mapping.get(field)
        if requested:
            if requested not in headers:
                raise ValueError(f"column for {field} not found: {requested}")
            result[field] = requested
            continue
        aliases = _DEFAULTS.get(field, (field,))
        found = next(
            (lower.get(alias.casefold()) for alias in aliases if alias.casefold() in lower), None
        )
        if found:
            result[field] = found
    if "base" in mapping:
        if mapping["base"] not in headers:
            raise ValueError(f"column for base not found: {mapping['base']}")
        result["frequency"] = mapping["base"]
    if "query" not in result:
        raise ValueError("import needs an explicit query mapping or a norm/phrase/query column")
    return result


def _parse(file, sheet=None, mapping=None):
    path = Path(file)
    if not path.is_file():
        raise FileNotFoundError(f"import file not found: {path}")
    if path.stat().st_size > MAX_IMPORT_BYTES:
        raise ValueError(f"import file exceeds the {MAX_IMPORT_BYTES}-byte limit")
    headers, source_rows = _read_table(path, sheet)
    applied = _mapping(headers, mapping)
    parsed, seen = [], set()
    for line, row in enumerate(source_rows, 2):
        raw = row.get(applied["query"])
        phrase = normalize(str(raw)) if raw is not None else ""
        if not phrase:
            continue
        values = {}
        for source_field, target_field in (
            ("frequency", "base"),
            ("quoted", "quoted"),
            ("exact", "exact"),
            ("impr", "impr"),
            ("pos_y", "pos_y"),
            ("pos_g", "pos_g"),
        ):
            column = applied.get(source_field)
            value = row.get(column) if column else None
            if value is None or str(value).strip().casefold() in {"", "null", "--", "none"}:
                continue
            if isinstance(value, bool) or (isinstance(value, float) and not value.is_integer()):
                raise ValueError(f"invalid {source_field} on input row {line}")
            try:
                number = int(value)
            except (TypeError, ValueError) as error:
                raise ValueError(f"invalid {source_field} on input row {line}") from error
            if number < 0:
                raise ValueError(f"negative {source_field} on input row {line}")
            values[target_field] = number
        group = str(row.get(applied["group"]) or "").strip() if applied.get("group") else ""
        subgroup = (
            str(row.get(applied["subgroup"]) or "").strip() if applied.get("subgroup") else ""
        )
        if group and not subgroup and ("/" in group or "\\" in group):
            parts = [part.strip() for part in group.replace("\\", "/").split("/") if part.strip()]
            group, subgroup = (parts[0], " / ".join(parts[1:])) if parts else ("", "")
        for field, column in (
            ("group_name", "group"),
            ("subgroup", "subgroup"),
            ("url", "url"),
            ("comment", "comment"),
        ):
            value = (
                group
                if field == "group_name"
                else subgroup
                if field == "subgroup"
                else row.get(applied.get(column))
                if applied.get(column)
                else None
            )
            if value is not None and str(value).strip():
                values[field] = str(value).strip()
        parsed.append({"query": phrase, "values": values, "duplicate_in_file": phrase in seen})
        seen.add(phrase)
    deduped = {}
    for item in parsed:
        if item["query"] not in deduped:
            deduped[item["query"]] = item
        else:
            deduped[item["query"]]["values"].update(item["values"])
    return list(deduped.values()), applied, len(parsed)


def plan(store, cfg, file, *, sheet=None, mapping=None, merge="add"):
    rows, applied, input_rows = _parse(file, sheet, mapping)
    return _plan_rows(
        store, file, sheet=sheet, merge=merge, rows=rows, applied=applied, input_rows=input_rows
    )


def _plan_rows(store, file, *, sheet, merge, rows, applied, input_rows):
    if merge not in {"add", "update", "replace"}:
        raise ValueError("merge policy must be add, update, or replace")
    if merge == "replace" and not rows:
        raise ValueError(
            "replace import contains no valid queries; refusing to clear the semantic core"
        )
    current = {r[0] for r in store.db.execute("SELECT norm FROM phrases")}
    existing = [item for item in rows if item["query"] in current]
    new = [item for item in rows if item["query"] not in current]
    groups: dict[str, int] = {}
    for item in rows:
        group = item["values"].get("group_name") or "(ungrouped)"
        groups[group] = groups.get(group, 0) + 1
    return {
        "contract": "seohead.semantics.operation.v1",
        "mode": "preview",
        "state": "complete",
        "format": "seohead.semantics.import-plan.v1",
        "file_type": Path(file).suffix.lower().lstrip("."),
        "sheet": sheet,
        "merge": merge,
        "applied_mapping": applied,
        "input_rows": input_rows,
        "unique_queries": len(rows),
        "new": len(new) if merge != "replace" else len(rows),
        "already_in_core": len(existing) if merge != "replace" else 0,
        "duplicates_in_input": input_rows - len(rows),
        "group_counts": groups,
        "preview": rows[:10],
        "replacement_rows": len(current) if merge == "replace" else None,
    }


def run(store, cfg, file, *, sheet=None, mapping=None, merge="add"):
    rows, applied, input_rows = _parse(file, sheet, mapping)
    preview = _plan_rows(
        store, file, sheet=sheet, merge=merge, rows=rows, applied=applied, input_rows=input_rows
    )
    if merge == "replace":
        store.db.execute("DELETE FROM edges")
        store.db.execute("DELETE FROM serp")
        store.db.execute("DELETE FROM domains")
        store.db.execute("DELETE FROM clusters")
        store.db.execute("DELETE FROM phrases")
        if store.db.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='landing_candidates'"
        ).fetchone():
            store.db.execute("DELETE FROM landing_candidates")
    added = updated = skipped = 0
    region = str((cfg.get("regions") or [225])[0])
    for item in rows:
        phrase, values = item["query"], item["values"]
        is_new = store.upsert(phrase, src="import", region=region)
        added += int(is_new)
        if not is_new and merge == "add":
            skipped += 1
            continue
        current = store.db.execute("SELECT * FROM phrases WHERE norm=?", (phrase,)).fetchone()
        changed = (
            values
            if merge in {"update", "replace"}
            else {key: value for key, value in values.items() if current[key] is None}
        )
        if changed:
            store.set_fields(phrase, **changed)
            updated += int(not is_new)
        if has_observed_demand(values.get("impr"), values.get("pos_y"), values.get("pos_g")):
            store.set_status(phrase, "kept", reason="observed search demand", stage="import")
    store.db.execute(
        "INSERT INTO semantic_audit(at,action,entity,before_json,after_json,evidence_json) "
        "VALUES(?,?,?,?,?,?)",
        (
            datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
            "import",
            "phrases",
            json.dumps(
                {"merge": "replace", "rows": preview["replacement_rows"]}, ensure_ascii=False
            )
            if merge == "replace"
            else None,
            json.dumps({"merge": merge, "rows": added + updated}, ensure_ascii=False),
            json.dumps({"file_type": preview["file_type"], "sheet": sheet}, ensure_ascii=False),
        ),
    )
    store.commit()
    print(
        f"import: rows={preview['unique_queries']} added={added} "
        f"updated={updated} skipped_existing={skipped}"
    )
    return {**preview, "added": added, "updated": updated, "skipped_existing": skipped}
